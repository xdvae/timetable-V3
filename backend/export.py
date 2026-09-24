"""Build a day x period grid for a given view and export it as XLSX/CSV/HTML."""
import csv
import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side


def build_lane_rows(scheduled_classes, days, period_labels, cell_html_func):
    """
    Turns a flat list of ScheduledClass rows into a day-by-day set of
    'lanes': each lane is one horizontal timeline for that day. A
    multi-period block (e.g. a 2-hour lab) becomes ONE cell spanning two
    columns (colspan) instead of repeating its text in every column it
    covers. If two sessions genuinely overlap in time for the same view
    (e.g. two different lab groups of the same section running parallel
    labs in different rooms), they're placed in separate lanes stacked
    under the same day, so it's visually obvious they're two different,
    simultaneous, non-conflicting sessions rather than one thing shown
    twice.
    """
    n = len(period_labels)
    by_day = {d: [] for d in days}
    for sc in scheduled_classes:
        if sc.day in by_day:
            by_day[sc.day].append(sc)

    day_rows = []
    for d in days:
        blocks = sorted(by_day[d], key=lambda sc: sc.start_period)
        lane_end = []
        lane_items = []
        for sc in blocks:
            placed = False
            for li in range(len(lane_end)):
                if sc.start_period >= lane_end[li]:
                    lane_items[li].append(sc)
                    lane_end[li] = sc.start_period + sc.length
                    placed = True
                    break
            if not placed:
                lane_items.append([sc])
                lane_end.append(sc.start_period + sc.length)
        if not lane_items:
            lane_items = [[]]

        lanes_cells = []
        for items in lane_items:
            occ = [None] * n
            for sc in items:
                for p in range(sc.start_period, sc.start_period + sc.length):
                    if p < n:
                        occ[p] = sc
            cells = []
            p = 0
            while p < n:
                sc = occ[p]
                if sc is None:
                    cells.append({"empty": True, "colspan": 1})
                    p += 1
                else:
                    length = min(sc.length, n - p)
                    cells.append({"empty": False, "colspan": length, "html": cell_html_func(sc)})
                    p += length
            lanes_cells.append(cells)
        day_rows.append({"day": d, "lanes": lanes_cells})
    return day_rows


def build_grid(scheduled_classes, days, period_labels, cell_text_func):
    """
    scheduled_classes: list of ScheduledClass rows (already filtered to one view)
    Returns grid[day][period_index] -> list of text lines (usually 1 entry, but
    could be several if something couldn't be avoided).
    """
    grid = {d: {p: [] for p in range(len(period_labels))} for d in days}
    for sc in scheduled_classes:
        for p in range(sc.start_period, sc.start_period + sc.length):
            if sc.day in grid and p in grid[sc.day]:
                grid[sc.day][p].append(cell_text_func(sc))
    return grid


def grid_to_rows(grid, days, period_labels):
    """Rows: header row + one row per day, columns = periods."""
    header = ["Day/Time"] + period_labels
    rows = [header]
    for d in days:
        row = [d]
        for p in range(len(period_labels)):
            cell = grid[d][p]
            row.append(" / ".join(cell) if cell else "")
        rows.append(row)
    return rows


def export_csv(grid, days, period_labels):
    rows = grid_to_rows(grid, days, period_labels)
    buf = io.StringIO()
    writer = csv.writer(buf)
    for r in rows:
        writer.writerow(r)
    return buf.getvalue()


def export_html(grid, days, period_labels, title="Timetable"):
    rows = grid_to_rows(grid, days, period_labels)
    html = [f"<html><head><meta charset='utf-8'><title>{title}</title>",
            "<style>",
            "body{font-family:Arial,sans-serif;padding:20px;}",
            "table{border-collapse:collapse;width:100%;margin-bottom:30px;}",
            "th,td{border:1px solid #999;padding:8px;text-align:center;font-size:13px;}",
            "th{background:#2c3e50;color:#fff;}",
            "td.day{background:#f0f0f0;font-weight:bold;}",
            "h2{font-family:Arial,sans-serif;}",
            "h3{font-family:Arial,sans-serif;color:#2c3e50;margin-top:30px;}",
            "@media print{ body{padding:0;} h3{page-break-before:auto;} }",
            "</style></head><body>",
            f"<h2>{title}</h2>",
            "<table>"]
    header = rows[0]
    html.append("<tr>" + "".join(f"<th>{h}</th>" for h in header) + "</tr>")
    for row in rows[1:]:
        html.append("<tr><td class='day'>" + row[0] + "</td>" +
                     "".join(f"<td>{c}</td>" for c in row[1:]) + "</tr>")
    html.append("</table></body></html>")
    return "\n".join(html)


def _draw_grid_on_sheet(ws, grid, days, period_labels, title):
    header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    day_fill = PatternFill(start_color="ECF0F1", end_color="ECF0F1", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(period_labels) + 1)

    header_row = 3
    ws.cell(row=header_row, column=1, value="Day / Time").font = header_font
    ws.cell(row=header_row, column=1).fill = header_fill
    ws.cell(row=header_row, column=1).border = border
    ws.cell(row=header_row, column=1).alignment = center
    for i, p in enumerate(period_labels):
        c = ws.cell(row=header_row, column=2 + i, value=p)
        c.font = header_font
        c.fill = header_fill
        c.border = border
        c.alignment = center

    for r, d in enumerate(days):
        row_idx = header_row + 1 + r
        dc = ws.cell(row=row_idx, column=1, value=d)
        dc.font = Font(bold=True)
        dc.fill = day_fill
        dc.border = border
        dc.alignment = center
        for i in range(len(period_labels)):
            cell_lines = grid[d][i]
            val = "\n".join(cell_lines) if cell_lines else ""
            c = ws.cell(row=row_idx, column=2 + i, value=val)
            c.border = border
            c.alignment = center

    ws.column_dimensions["A"].width = 14
    for i in range(len(period_labels)):
        col_letter = ws.cell(row=header_row, column=2 + i).column_letter
        ws.column_dimensions[col_letter].width = 22
    for r in range(len(days)):
        ws.row_dimensions[header_row + 1 + r].height = 40


def export_xlsx(grid, days, period_labels, title="Timetable"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Timetable"
    _draw_grid_on_sheet(ws, grid, days, period_labels, title)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
